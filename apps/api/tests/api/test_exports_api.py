"""Tests for export API endpoints (RF-044, RF-045, RF-046)."""

from __future__ import annotations

import datetime
import io
import uuid

import pytest
from openpyxl import load_workbook

from app.core.models_base import utcnow
from app.modules.search.models import ExportJobKind, ExportJobModel, ExportJobStatus
from tests.api.conftest import LoginAs, SeededApi


@pytest.fixture
def fake_storage(seeded_api: SeededApi, monkeypatch: pytest.MonkeyPatch) -> dict[str, bytes]:
    """In-memory object storage mock for API tests."""
    objects: dict[str, bytes] = {}

    def put_bytes(key: str, data: bytes, content_type: str) -> None:
        objects[key] = data

    def presign_get(key: str, ttl_seconds: int = 300) -> str:
        return f"https://s3.local/{key}?signature=dummy"

    monkeypatch.setattr(seeded_api.app.state.storage, "put_bytes", put_bytes)
    monkeypatch.setattr(seeded_api.app.state.storage, "presign_get", presign_get)
    return objects


def test_export_excel_sync_for_cataloguer(seeded_api: SeededApi) -> None:
    """Cataloguer exports Excel <= 5000 rows synchronously; sensitive fields omitted."""
    client = seeded_api.login_as("CATALOGUER")
    response = client.post(
        "/api/v1/reports/export-excel",
        json={"filters": {}},
    )
    assert response.status_code == 200
    assert "application/vnd.openxmlformats" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]

    # Verify sensitive columns omitted for cataloguer
    omitted = response.headers.get("x-omited-columns", "")
    assert "lender_name" in omitted or "valuation" in omitted

    # Verify valid xlsx workbook
    wb = load_workbook(io.BytesIO(response.content))
    assert "Resultados" in wb.sheetnames
    assert "Filtros" in wb.sheetnames


def test_full_export_permission_and_deduplication(
    login_as: LoginAs, seeded_api: SeededApi, fake_storage: dict[str, bytes]
) -> None:
    """Full export is restricted to Admin; calling it while in progress returns existing job."""
    cataloguer = login_as("CATALOGUER")
    resp_cat = cataloguer.get("/api/v1/exports/full")
    assert resp_cat.status_code == 403

    admin = login_as("ADMIN")
    resp_admin = admin.get("/api/v1/exports/full")
    assert resp_admin.status_code in (200, 202)
    job_data = resp_admin.json()
    assert job_data["id"] is not None
    assert job_data["status"] in ("PENDING", "RUNNING", "READY")
    job_id = job_data["id"]

    # Deduplication: requesting again while job is running returns same job
    resp_repeat = admin.get("/api/v1/exports/full")
    assert resp_repeat.status_code in (200, 202)
    assert resp_repeat.json()["id"] == job_id


def test_export_job_status_and_download_flow(
    login_as: LoginAs, seeded_api: SeededApi, fake_storage: dict[str, bytes]
) -> None:
    """Export job polling, status, and download-url behavior."""
    admin = login_as("ADMIN")
    admin_id = uuid.UUID(admin.get("/api/v1/auth/me").json()["id"])

    # Create a job in DB
    with seeded_api.session() as session:
        job = ExportJobModel(
            kind=ExportJobKind.RESULTS,
            requested_by_id=admin_id,
            format="xlsx",
            params={},
            status=ExportJobStatus.RUNNING,
        )
        session.add(job)
        session.commit()
        job_id = str(job.id)

    # 1. While running, download-url returns 409
    resp_dl = admin.get(f"/api/v1/exports/jobs/{job_id}/download-url")
    assert resp_dl.status_code == 409

    # 2. Update job to READY
    with seeded_api.session() as session:
        j = session.get(ExportJobModel, uuid.UUID(job_id))
        j.status = ExportJobStatus.READY
        j.storage_key = f"exports/{job_id}/results.xlsx"
        j.sha256 = "dummy_sha256"
        j.row_count = 10
        j.expires_at = utcnow() + datetime.timedelta(hours=24)
        session.commit()

    resp_ready = admin.get(f"/api/v1/exports/jobs/{job_id}")
    assert resp_ready.status_code == 200
    assert resp_ready.json()["status"] == "READY"
    assert resp_ready.json()["row_count"] == 10

    resp_dl_ready = admin.get(f"/api/v1/exports/jobs/{job_id}/download-url")
    assert resp_dl_ready.status_code == 200
    assert "download_url" in resp_dl_ready.json()

    # 3. If job is expired, download-url returns 410
    with seeded_api.session() as session:
        j = session.get(ExportJobModel, uuid.UUID(job_id))
        j.expires_at = utcnow() - datetime.timedelta(hours=1)
        session.commit()

    resp_expired = admin.get(f"/api/v1/exports/jobs/{job_id}/download-url")
    assert resp_expired.status_code == 410
