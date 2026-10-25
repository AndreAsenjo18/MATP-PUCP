"""Spec importacion-datos — decisiones, aprobación, aplicación, bitácora e imágenes por HTTP.

Change importacion-pipeline-reconciliacion, tareas 4.1 (Req: Decisiones por fila y por campo
con resolución de conflictos), 4.2 (Req: Aplicación atómica con datos de origen y reversión del
lote; RF-027, RF-008), 4.3 (Req: Bitácora descargable del lote; RF-028) y 5.1 (Req: Revisión
manual de imágenes extraídas de Excel; RF-029).

Every test builds its own batch with a unique tag in the denominations, so the pieces an
applied batch creates never turn the rows of another test into possible duplicates.
"""

import io
import uuid
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from sqlalchemy import select

from app.modules.audit.context import AuditContext, audit_context
from app.modules.audit.models import AuditLog, AuditOrigin
from app.modules.catalog.enums import LEGAL_OWNER_PUCP, TenureRegime
from app.modules.catalog.models import Piece, PieceSourceRecord
from app.modules.catalog.service import change_tenure, create_piece, update_piece
from app.modules.identification.models import PieceIdentifier
from app.modules.identification.normalization import TYPE_COLLECTION, TYPE_INVENTORY
from app.modules.identification.service import add_identifier
from app.modules.imports.models import ImportBatch, ImportImage, ImportRow
from app.modules.media.models import MediaAsset
from app.modules.users.models import AppUser
from tests.api.conftest import ADMIN, Pipeline
from tests.api.test_import_pipeline_api import (
    BASE,
    SABANA,
    Catalog,
    _catalog,
    diffs,
    get_batch,
    put_mapping,
    upload,
    validate,
)

SHEET_MAPPING = {
    "columns": [
        {"source_column": "DENOMINACIÓN", "target_field": "denomination"},
        {"source_column": "CÓDIGOS", "target_field": "identifier", "split_compound": True},
        {"source_column": "RÉGIMEN", "target_field": "tenure_regime"},
        {"source_column": "PROCEDENCIA", "target_field": "provenance"},
        {"source_column": "AUTOR", "target_field": "author"},
    ]
}


def _tag() -> str:
    return uuid.uuid4().hex[:8]


def _sheet(catalog: Catalog, tag: str) -> bytes:
    """One row per classification rule (design D4), with values unique to ``tag``.

    Row 2 updates the target piece, 3 is new, 4 has codes of two pieces, 5 and 6 share a code I
    and 7 is a loan for use with a code I (RN-003).
    """
    code = int(tag, 16) % 90000 + 10000
    lines = [
        "DENOMINACIÓN;CÓDIGOS;RÉGIMEN;PROCEDENCIA;AUTOR;OBS. CONSULTORÍA",
        f"{catalog.target.denomination};{catalog.target_code};Propiedad;Procedencia {tag};;"
        f"Revisado {tag}",
        f"Pieza nueva {tag};S/N;Propiedad;;;Nueva {tag}",
        f"Contradictoria {tag};{catalog.third_code} / {catalog.other_collection_code};Propiedad;;;",
        f"Primera repetida {tag};I-{code};Propiedad;;;",
        f"Segunda repetida {tag};I {code};Propiedad;;;",
        f"Comodato {tag};I-{code + 1};Comodato;;;",
    ]
    return ("\n".join(lines) + "\n").encode("utf-8")


@dataclass
class Batch:
    id: str
    tag: str
    catalog: Catalog
    rows: dict[int, dict[str, Any]]


def _rows(pipeline: Pipeline, batch_id: str) -> dict[int, dict[str, Any]]:
    page = diffs(pipeline.client, batch_id, page_size=100).json()
    return {row["source_row_number"]: row for row in page["items"]}


def prepared(pipeline: Pipeline) -> Batch:
    with pipeline.session() as session:
        catalog = _catalog(session)
        session.expunge_all()
    tag = _tag()
    batch_id = upload(pipeline.client, f"decisiones-{tag}.csv", _sheet(catalog, tag))["batch_id"]
    assert put_mapping(pipeline.client, batch_id, SHEET_MAPPING).status_code == 200
    assert validate(pipeline.client, batch_id).status_code == 202
    return Batch(batch_id, tag, catalog, _rows(pipeline, batch_id))


def decide(
    pipeline: Pipeline, batch: Batch, row: int, role: str = "CATALOGUER", **body: Any
) -> Any:
    row_id = batch.rows[row]["id"]
    return pipeline.login(role).patch(f"{BASE}/{batch.id}/rows/{row_id}", json=body)


def summary(pipeline: Pipeline, batch_id: str) -> dict[str, Any]:
    approval = get_batch(pipeline.client, batch_id)["approval_summary"]
    assert approval is not None
    return approval


def confirm(
    pipeline: Pipeline, batch_id: str, token: str, role: str = "COLLECTIONS_MANAGER"
) -> Any:
    return pipeline.login(role).post(
        f"{BASE}/{batch_id}/confirm",
        json={"summary_token": token, "comment": "Revisado (sintético)"},
    )


def _piece_of(row: dict[str, Any], identifier_type: str) -> str:
    return next(
        match["piece_id"]
        for match in row["matches"]
        if match["kind"] == "piece"
        and any(item["identifier_type"] == identifier_type for item in match["identifiers"])
    )


# --- 4.1 Decisiones por fila y por campo --------------------------------------------------------
def test_resolving_a_conflict_choosing_the_target_piece(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)
    piece_a = _piece_of(batch.rows[4], TYPE_INVENTORY)

    response = decide(pipeline, batch, 4, decision="ACCEPTED", target_piece_id=piece_a)

    assert response.status_code == 200, response.text
    row = response.json()
    assert row["classification"] == "CONFLICT"  # the computed classification is kept
    assert row["action"] == "UPDATE" and row["target_piece_id"] == piece_a
    assert {change["field"] for change in row["diff"]} >= {"denomination"}
    assert row["resolution"]["target_piece_id"] == piece_a
    log = pipeline.client.get(f"{BASE}/{batch.id}/log").json()
    entry = next(entry for entry in log["entries"] if entry["source_row_number"] == 4)
    assert entry["resolution"]["target_piece_id"] == piece_a


def test_accepting_a_conflict_without_a_target_is_422(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    response = decide(pipeline, batch, 5, decision="ACCEPTED")

    assert response.status_code == 422
    assert response.json()["code"] == "conflict_unresolved"
    assert _rows(pipeline, batch.id)[5]["decision"] == "PENDING"


def test_rejecting_requires_a_reason(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    missing = decide(pipeline, batch, 3, decision="REJECTED", reason="  ")
    given = decide(pipeline, batch, 3, decision="REJECTED", reason="Pieza fuera del alcance")

    assert missing.status_code == 422 and missing.json()["code"] == "reason_required"
    assert given.status_code == 200
    assert given.json()["decision"] == "REJECTED" and given.json()["action"] == "SKIP"
    assert given.json()["decision_reason"] == "Pieza fuera del alcance"


def test_row_with_validation_errors_cannot_be_accepted(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    response = decide(pipeline, batch, 7, decision="ACCEPTED")

    assert response.status_code == 422
    assert response.json()["code"] == "row_has_errors"


def test_fields_of_an_update_are_accepted_or_excluded(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    unknown = decide(pipeline, batch, 2, decision="ACCEPTED", field_decisions={"x": "accept"})
    response = decide(
        pipeline,
        batch,
        2,
        decision="ACCEPTED",
        field_decisions={"provenance": "exclude", "author": "accept"},
    )

    assert unknown.status_code == 422 and unknown.json()["code"] == "unknown_diff_field"
    assert response.status_code == 200
    assert response.json()["resolution"]["field_decisions"] == {
        "provenance": "exclude",
        "author": "accept",
    }


def test_a_target_with_another_code_i_is_rejected(pipeline: Pipeline) -> None:
    """RN-002: resolving to a piece never changes its code I."""
    batch = prepared(pipeline)

    response = decide(
        pipeline, batch, 5, decision="ACCEPTED", target_piece_id=str(batch.catalog.target.id)
    )

    assert response.status_code == 422
    assert response.json()["code"] == "inventory_code_mismatch"


def test_a_row_whose_code_i_is_in_use_cannot_be_new(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    response = decide(pipeline, batch, 4, decision="ACCEPTED", as_new=True)

    assert response.status_code == 422
    assert response.json()["code"] == "inventory_code_in_use"


# --- 4.2 Aprobación y aplicación atómica --------------------------------------------------------
def test_cataloguer_cannot_approve(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)

    response = confirm(
        pipeline, batch.id, summary(pipeline, batch.id)["summary_token"], "CATALOGUER"
    )

    assert response.status_code == 403


def test_approving_a_mapped_batch_is_409(pipeline: Pipeline) -> None:
    batch_id = upload(pipeline.client, "sabana.xlsx", SABANA)["batch_id"]
    put_mapping(
        pipeline.client,
        batch_id,
        {"columns": [{"source_column": "DENOMINACIÓN", "target_field": "denomination"}]},
    )

    response = confirm(pipeline, batch_id, "0" * 32)

    assert response.status_code == 409
    body = response.json()
    assert body["code"] == "invalid_batch_transition"
    assert body["details"]["current_status"] == "MAPPED"
    assert "mapear columnas" in body["message"]


def test_a_stale_summary_is_409_with_the_current_one(pipeline: Pipeline) -> None:
    batch = prepared(pipeline)
    shown = summary(pipeline, batch.id)
    assert decide(pipeline, batch, 3, decision="EXCLUDED").status_code == 200

    response = confirm(pipeline, batch.id, shown["summary_token"])

    assert response.status_code == 409
    body = response.json()
    assert body["code"] == "summary_changed"
    assert body["details"]["summary"]["create"] == shown["create"] - 1
    assert get_batch(pipeline.client, batch.id)["status"] == "IN_PREVIEW"


@pytest.fixture(scope="module")
def applied(pipeline: Pipeline) -> Batch:
    """A controlled batch approved with row 4 rejected and the other conflicts left pending."""
    batch = prepared(pipeline)
    rejected = decide(pipeline, batch, 4, decision="REJECTED", reason="Códigos de dos piezas")
    assert rejected.status_code == 200
    shown = summary(pipeline, batch.id)
    assert shown["create"] == 1 and shown["update"] == 1 and shown["skip"] == 4

    response = confirm(pipeline, batch.id, shown["summary_token"])

    assert response.status_code == 200, response.text
    assert response.json()["status"] == "APPROVED"  # applied in the background
    return batch


def test_application_creates_and_updates_pieces(pipeline: Pipeline, applied: Batch) -> None:
    batch = get_batch(pipeline.client, applied.id)
    assert batch["status"] == "APPLIED" and batch["processing"] is False
    assert batch["counts"]["APPLIED_CREATED"] == 1
    assert batch["counts"]["APPLIED_UPDATED"] == 1
    assert batch["approved_by_id"] is not None and batch["approval_summary"] is None
    with pipeline.session() as session:
        created = session.scalar(
            select(Piece).where(Piece.denomination == f"Pieza nueva {applied.tag}")
        )
        assert created is not None and created.tenure_regime == TenureRegime.OWNED
        target = session.get(Piece, applied.catalog.target.id)
        assert target is not None
        assert target.provenance == f"Procedencia {applied.tag}"
        assert target.author == applied.catalog.target.author  # empty cell → keep [B10]
        row = session.scalar(
            select(ImportRow).where(
                ImportRow.batch_id == uuid.UUID(applied.id), ImportRow.source_row_number == 3
            )
        )
        assert row is not None and row.target_piece_id == created.id


def test_every_applied_piece_keeps_its_source_data(pipeline: Pipeline, applied: Batch) -> None:
    """RF-008: unmapped columns with the row number and the batch."""
    with pipeline.session() as session:
        records = session.scalars(
            select(PieceSourceRecord).where(
                PieceSourceRecord.import_batch_id == uuid.UUID(applied.id)
            )
        ).all()
    by_row = {record.source_row_number: record for record in records}
    assert set(by_row) == {2, 3}
    assert by_row[2].piece_id == applied.catalog.target.id
    assert by_row[2].payload == {"OBS. CONSULTORÍA": f"Revisado {applied.tag}"}
    assert by_row[3].source_file_name == f"decisiones-{applied.tag}.csv"


def test_application_is_audited_as_one_import_change_set(
    pipeline: Pipeline, applied: Batch
) -> None:
    with pipeline.session() as session:
        batch = session.get(ImportBatch, uuid.UUID(applied.id))
        assert batch is not None and batch.change_set_id is not None
        entries = session.scalars(
            select(AuditLog).where(AuditLog.change_set_id == batch.change_set_id)
        ).all()
    assert {entry.origin for entry in entries} == {AuditOrigin.IMPORT}
    assert {entry.origin_ref for entry in entries} == {applied.id}
    changed = {(entry.entity_id, entry.field) for entry in entries}
    assert (applied.catalog.target.id, "provenance") in changed


def _user_id(session: Any, email: str) -> uuid.UUID:
    user = session.scalar(select(AppUser).where(AppUser.email == email))
    assert user is not None
    return user.id


def test_a_rule_violated_at_apply_rolls_back_every_row(pipeline: Pipeline) -> None:
    """The target becomes a loan for use after validation: its new code I breaks RN-003."""
    tag = _tag()
    code = int(tag, 16) % 9000 + 1000
    with pipeline.session() as session:
        context = AuditContext(origin=AuditOrigin.MANUAL, user_id=_user_id(session, ADMIN))
        with audit_context(session, context):
            target = create_piece(
                session, denomination=f"Pasará a comodato {tag}", tenure_regime=TenureRegime.OWNED
            )
            add_identifier(session, target, TYPE_COLLECTION, f"MMZ {code}")
        session.commit()
        target_id = target.id
    content = (
        "DENOMINACIÓN;CÓDIGOS;RÉGIMEN\n"
        f"Pasará a comodato {tag};MMZ {code} / I-9{code};Propiedad\n"
        f"Otra nueva {tag};S/N;Propiedad\n"
    ).encode()
    batch_id = upload(pipeline.client, f"comodato-{tag}.csv", content)["batch_id"]
    put_mapping(pipeline.client, batch_id, {"columns": SHEET_MAPPING["columns"][:3]})
    validate(pipeline.client, batch_id)
    rows = _rows(pipeline, batch_id)
    assert rows[2]["classification"] == "UPDATE"
    assert f"identifier:{TYPE_INVENTORY}" in {change["field"] for change in rows[2]["diff"]}
    with pipeline.session() as session:
        context = AuditContext(origin=AuditOrigin.MANUAL, user_id=_user_id(session, ADMIN))
        with audit_context(session, context):
            piece = session.get(Piece, target_id)
            assert piece is not None
            change_tenure(session, piece, TenureRegime.LOAN_FOR_USE, "Contrato sintético")
        session.commit()

    assert (
        confirm(pipeline, batch_id, summary(pipeline, batch_id)["summary_token"]).status_code == 200
    )

    failed = get_batch(pipeline.client, batch_id)
    assert failed["status"] == "FAILED_APPLY"
    assert "Fila 2" in failed["status_reason"] and "comodato" in failed["status_reason"]
    with pipeline.session() as session:
        assert (
            session.scalar(select(Piece).where(Piece.denomination == f"Otra nueva {tag}")) is None
        )
        codes = session.scalars(
            select(PieceIdentifier).where(
                PieceIdentifier.piece_id == target_id,
                PieceIdentifier.identifier_type_code == TYPE_INVENTORY,
            )
        ).all()
        assert codes == []

    # Correct and retry: excluding the row brings the batch back to preview (design D1).
    row_id = rows[2]["id"]
    excluded = pipeline.client.patch(
        f"{BASE}/{batch_id}/rows/{row_id}",
        json={"decision": "EXCLUDED"},
    )
    assert excluded.status_code == 200
    assert get_batch(pipeline.client, batch_id)["status"] == "IN_PREVIEW"
    assert (
        confirm(pipeline, batch_id, summary(pipeline, batch_id)["summary_token"]).status_code == 200
    )
    assert get_batch(pipeline.client, batch_id)["status"] == "APPLIED"
    with pipeline.session() as session:
        assert session.scalar(select(Piece).where(Piece.denomination == f"Otra nueva {tag}"))


def test_a_batch_upgrades_a_loan_for_use_with_the_batch_as_document(pipeline: Pipeline) -> None:
    """Comodato -> propiedad: the application supplies the batch as supporting document (D4)."""
    tag = _tag()
    code = int(tag, 16) % 9000 + 1000
    with pipeline.session() as session:
        context = AuditContext(origin=AuditOrigin.MANUAL, user_id=_user_id(session, ADMIN))
        with audit_context(session, context):
            target = create_piece(
                session,
                denomination=f"Vuelve a propiedad {tag}",
                tenure_regime=TenureRegime.LOAN_FOR_USE,
            )
            add_identifier(session, target, TYPE_COLLECTION, f"MMZ {code}")
        session.commit()
        target_id = target.id
    content = (
        f"DENOMINACIÓN;CÓDIGOS;RÉGIMEN\nVuelve a propiedad {tag};MMZ {code};Propiedad\n"
    ).encode()
    batch_id = upload(pipeline.client, f"comodato-{tag}.csv", content)["batch_id"]
    put_mapping(pipeline.client, batch_id, {"columns": SHEET_MAPPING["columns"][:3]})
    validate(pipeline.client, batch_id)
    rows = _rows(pipeline, batch_id)
    assert rows[2]["classification"] == "UPDATE"
    assert "tenure_regime" in {change["field"] for change in rows[2]["diff"]}

    response = confirm(pipeline, batch_id, summary(pipeline, batch_id)["summary_token"])

    assert response.status_code == 200
    applied = get_batch(pipeline.client, batch_id)
    assert applied["status"] == "APPLIED"
    assert applied["counts"]["APPLIED_UPDATED"] == 1
    with pipeline.session() as session:
        piece = session.get(Piece, target_id)
        assert piece is not None
        assert piece.tenure_regime == TenureRegime.OWNED
        assert piece.legal_owner == LEGAL_OWNER_PUCP
        entry = session.scalar(
            select(AuditLog).where(
                AuditLog.origin_ref == batch_id,
                AuditLog.entity_id == target_id,
                AuditLog.field == "tenure_regime",
            )
        )
        assert entry is not None
        assert f"lote de importación {batch_id}" in (entry.reason or "")


# --- 4.3 Bitácora -------------------------------------------------------------------------------
def test_log_of_an_applied_batch_lists_the_rows_not_applied(
    pipeline: Pipeline, applied: Batch
) -> None:
    response = pipeline.client.get(f"{BASE}/{applied.id}/log")

    assert response.status_code == 200
    log = response.json()
    assert log["applied_rows"] == 2 and log["rejected_rows"] == 4
    assert log["batch"]["approved_by"] is not None and log["batch"]["uploaded_by"] is not None
    assert {"UPLOADED", "IN_PREVIEW", "APPROVED", "APPLIED"} <= set(
        log["batch"]["stage_timestamps"]
    )
    reasons = {entry["source_row_number"]: entry["reason"] for entry in log["entries"]}
    assert set(reasons) == {4, 5, 6, 7}
    assert reasons[4] == "Códigos de dos piezas"
    assert reasons[5].startswith("Conflicto sin resolver")
    assert "RN-003" in reasons[7]
    assert log["download_url"] == f"/api/v1/imports/{applied.id}/log?format=xlsx"


def test_log_downloads_as_a_spreadsheet(pipeline: Pipeline, applied: Batch) -> None:
    response = pipeline.client.get(f"{BASE}/{applied.id}/log", params={"format": "xlsx"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Resumen", "Filas no aplicadas", "Resoluciones"]
    summary_values = {row[0]: row[1] for row in workbook["Resumen"].iter_rows(values_only=True)}
    assert summary_values["Archivo"] == f"decisiones-{applied.tag}.csv"
    assert summary_values["Filas aplicadas"] == 2
    skipped = list(workbook["Filas no aplicadas"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in skipped] == [4, 5, 6, 7]
    assert skipped[0][4] == "Códigos de dos piezas"


def test_log_of_a_batch_failed_in_ingestion_shows_the_reason(pipeline: Pipeline) -> None:
    batch_id = upload(pipeline.client, "inventario.pdf", b"%PDF-1.4 sintetico")["batch_id"]

    log = pipeline.client.get(f"{BASE}/{batch_id}/log").json()

    assert log["batch"]["status"] == "FAILED_INGESTION"
    assert ".pdf" in log["batch"]["status_reason"]
    assert log["applied_rows"] is None and log["rejected_rows"] is None
    assert log["entries"] == [] and log["batch"]["counts"] == {}


# --- 5.1 Imágenes extraídas del Excel -----------------------------------------------------------
def test_sabana_images_are_proposed_for_their_rows_or_left_unassigned(pipeline: Pipeline) -> None:
    batch_id = upload(pipeline.client, "sabana.xlsx", SABANA)["batch_id"]
    put_mapping(
        pipeline.client,
        batch_id,
        {
            "columns": [{"source_column": "DENOMINACIÓN", "target_field": "denomination"}],
            "defaults": {"tenure_regime": "Propiedad"},
        },
    )
    validate(pipeline.client, batch_id)

    page = diffs(pipeline.client, batch_id, page_size=100).json()

    by_row = {row["source_row_number"]: row["images"] for row in page["items"]}
    assert [image["image_number"] for image in by_row[2]] == [1]
    assert [image["image_number"] for image in by_row[4]] == [2]
    assert sum(len(images) for images in by_row.values()) == 2
    tray = page["unassigned_images"]
    assert [(image["image_number"], image["anchor_row_number"]) for image in tray] == [(3, 50)]
    assert {image["status"] for image in [*by_row[2], *tray]} == {"PENDING"}
    assert get_batch(pipeline.client, batch_id)["counts"]["images"] == 3

    image_id = by_row[2][0]["id"]
    early = pipeline.client.post(f"{BASE}/{batch_id}/images/{image_id}/associate", json={})
    assert early.status_code == 409 and early.json()["code"] == "batch_not_applied"


def _workbook_with_images(tag: str) -> bytes:
    """Two data rows; one image in row 2 and one in row 10, where there is no data."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "INVENTARIO"
    sheet.append(["DENOMINACIÓN", "CÓDIGOS"])
    sheet.append([f"Pieza con foto {tag}", "S/N"])
    sheet.append([f"Pieza sin foto {tag}", "S/N"])
    for anchor, color in (("D2", "red"), ("D10", "blue")):
        picture = io.BytesIO()
        Image.new("RGB", (40, 30), color).save(picture, format="JPEG")
        picture.seek(0)
        sheet.add_image(SheetImage(picture), anchor)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@pytest.fixture
def applied_with_images(pipeline: Pipeline) -> Iterator[tuple[str, str, dict[str, Any]]]:
    tag = _tag()
    batch_id = upload(pipeline.client, f"fotos-{tag}.xlsx", _workbook_with_images(tag))["batch_id"]
    put_mapping(
        pipeline.client,
        batch_id,
        {
            "columns": [
                {"source_column": "DENOMINACIÓN", "target_field": "denomination"},
                {"source_column": "CÓDIGOS", "target_field": "identifier", "split_compound": True},
            ],
            "defaults": {"tenure_regime": "Propiedad"},
        },
    )
    validate(pipeline.client, batch_id)
    page = diffs(pipeline.client, batch_id, page_size=100).json()
    assert (
        confirm(pipeline, batch_id, summary(pipeline, batch_id)["summary_token"]).status_code == 200
    )
    assert get_batch(pipeline.client, batch_id)["status"] == "APPLIED"
    yield batch_id, tag, page


def test_associating_an_image_registers_it_as_a_photo_of_the_piece(
    pipeline: Pipeline, applied_with_images: tuple[str, str, dict[str, Any]]
) -> None:
    batch_id, tag, page = applied_with_images
    row = next(row for row in page["items"] if row["source_row_number"] == 2)
    image_id = row["images"][0]["id"]
    url = f"{BASE}/{batch_id}/images/{image_id}/associate"

    response = pipeline.client.post(url, json={})
    again = pipeline.client.post(url, json={})

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "ASSOCIATED" and body["media_asset_id"] is not None
    assert again.status_code == 409 and again.json()["code"] == "image_already_associated"
    with pipeline.session() as session:
        piece = session.scalar(select(Piece).where(Piece.denomination == f"Pieza con foto {tag}"))
        asset = session.get(MediaAsset, uuid.UUID(body["media_asset_id"]))
        assert piece is not None and asset is not None
        assert asset.piece_id == piece.id and asset.is_primary
        assert asset.content_type == "image/jpeg" and asset.width_px == 40
        assert asset.extra_metadata is not None
        assert asset.extra_metadata["import_batch_id"] == batch_id
        assert asset.extra_metadata["source_row_number"] == 2
        image = session.get(ImportImage, uuid.UUID(image_id))
        assert image is not None
        assert asset.content_sha256 == image.content_sha256
    assert asset.storage_key in pipeline.stored
    assert f"pieces/{piece.id}/{asset.id}/thumb.webp" in pipeline.stored


def test_an_unassigned_image_needs_the_piece(
    pipeline: Pipeline, applied_with_images: tuple[str, str, dict[str, Any]]
) -> None:
    batch_id, tag, page = applied_with_images
    tray = page["unassigned_images"]
    assert [image["anchor_row_number"] for image in tray] == [10]
    url = f"{BASE}/{batch_id}/images/{tray[0]['id']}/associate"

    missing = pipeline.client.post(url, json={})
    with pipeline.session() as session:
        piece = session.scalar(select(Piece).where(Piece.denomination == f"Pieza sin foto {tag}"))
        assert piece is not None
        piece_id = str(piece.id)
    chosen = pipeline.client.post(url, json={"piece_id": piece_id})

    assert missing.status_code == 422 and missing.json()["code"] == "piece_required"
    assert chosen.status_code == 200 and chosen.json()["status"] == "ASSOCIATED"


# --- 5.2 Reversión del lote ----------------------------------------------------------------------
@dataclass
class AppliedBatch:
    id: str
    tag: str
    target_id: uuid.UUID
    reason: str = "Lote cargado con la hoja equivocada"


def _applied_update_and_create(pipeline: Pipeline) -> AppliedBatch:
    """A batch of two rows: row 2 updates a piece of its own, row 3 creates a new one."""
    tag = _tag()
    code = int(tag, 16) % 900000 + 100000
    with pipeline.session() as session:
        context = AuditContext(origin=AuditOrigin.MANUAL, user_id=_user_id(session, ADMIN))
        with audit_context(session, context):
            target = create_piece(
                session,
                denomination=f"Por actualizar {tag}",
                tenure_regime=TenureRegime.OWNED,
                provenance=f"Procedencia original {tag}",
            )
            add_identifier(session, target, TYPE_COLLECTION, f"MMZ {code}")
        session.commit()
        target_id = target.id
    content = (
        "DENOMINACIÓN;CÓDIGOS;RÉGIMEN;PROCEDENCIA\n"
        f"Por actualizar {tag};MMZ {code};Propiedad;Procedencia del lote {tag}\n"
        f"Creada por el lote {tag};S/N;Propiedad;\n"
    ).encode()
    batch_id = upload(pipeline.client, f"reversion-{tag}.csv", content)["batch_id"]
    mapping = {"columns": SHEET_MAPPING["columns"][:4]}
    assert put_mapping(pipeline.client, batch_id, mapping).status_code == 200
    assert validate(pipeline.client, batch_id).status_code == 202
    rows = _rows(pipeline, batch_id)
    assert rows[2]["classification"] == "UPDATE" and rows[3]["classification"] == "NEW"
    response = confirm(pipeline, batch_id, summary(pipeline, batch_id)["summary_token"])
    assert response.status_code == 200, response.text
    assert get_batch(pipeline.client, batch_id)["status"] == "APPLIED"
    return AppliedBatch(batch_id, tag, target_id)


def rollback(pipeline: Pipeline, batch: AppliedBatch, role: str = "ADMIN", **body: Any) -> Any:
    return pipeline.login(role).post(
        f"{BASE}/{batch.id}/rollback", json={"reason": batch.reason, **body}
    )


def test_rolling_back_a_batch_undoes_it_without_deleting_history(pipeline: Pipeline) -> None:
    """Scenario «Reversión del lote»: created pieces deleted logically, values restored."""
    batch = _applied_update_and_create(pipeline)

    response = rollback(pipeline, batch)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "REVERTED" and body["status_reason"] == batch.reason
    with pipeline.session() as session:
        target = session.get(Piece, batch.target_id)
        assert target is not None and target.deleted_at is None
        assert target.provenance == f"Procedencia original {batch.tag}"
        created = session.scalar(
            select(Piece)
            .where(Piece.denomination == f"Creada por el lote {batch.tag}")
            .execution_options(include_deleted=True)
        )
        assert created is not None
        assert created.deleted_at is not None and created.deletion_reason == batch.reason
        # The source data of each applied row stays as history (RN-005, append-only).
        records = session.scalars(
            select(PieceSourceRecord).where(
                PieceSourceRecord.import_batch_id == uuid.UUID(batch.id)
            )
        ).all()
        assert {record.source_row_number for record in records} == {2, 3}
        stored = session.get(ImportBatch, uuid.UUID(batch.id))
        assert stored is not None and stored.change_set_id is not None
        reversal = session.scalars(
            select(AuditLog).where(AuditLog.origin_ref == str(stored.change_set_id))
        ).all()
    assert reversal and {entry.reason for entry in reversal} == {batch.reason}


def test_a_later_edit_must_be_confirmed_before_rolling_back(pipeline: Pipeline) -> None:
    """RNF-007: an edit made after applying is never overwritten silently (design D7)."""
    batch = _applied_update_and_create(pipeline)
    with pipeline.session() as session:
        context = AuditContext(origin=AuditOrigin.MANUAL, user_id=_user_id(session, ADMIN))
        with audit_context(session, context):
            piece = session.get(Piece, batch.target_id)
            assert piece is not None
            update_piece(session, piece, provenance=f"Corregida a mano {batch.tag}")
            session.flush()
        session.commit()

    unconfirmed = rollback(pipeline, batch)

    assert unconfirmed.status_code == 409
    assert unconfirmed.json()["code"] == "revert_has_conflicts"
    conflicts = unconfirmed.json()["details"]["conflicts"]
    assert {(c["entity_id"], c["field"]) for c in conflicts} == {
        (str(batch.target_id), "provenance")
    }
    assert get_batch(pipeline.client, batch.id)["status"] == "APPLIED"

    confirmed = rollback(pipeline, batch, confirm_conflicts=True)

    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "REVERTED"
    with pipeline.session() as session:
        piece = session.get(Piece, batch.target_id)
        assert piece is not None and piece.provenance == f"Procedencia original {batch.tag}"


def test_only_an_administrator_rolls_back_a_batch(pipeline: Pipeline) -> None:
    batch = _applied_update_and_create(pipeline)

    response = rollback(pipeline, batch, role="COLLECTIONS_MANAGER")

    assert response.status_code == 403
    assert get_batch(pipeline.client, batch.id)["status"] == "APPLIED"


def test_a_batch_not_applied_or_already_reverted_cannot_be_rolled_back(
    pipeline: Pipeline,
) -> None:
    pending = prepared(pipeline)
    in_preview = pipeline.login("ADMIN").post(
        f"{BASE}/{pending.id}/rollback", json={"reason": "No está aplicado"}
    )
    batch = _applied_update_and_create(pipeline)
    assert rollback(pipeline, batch).status_code == 200

    again = rollback(pipeline, batch)

    assert in_preview.status_code == 409
    assert in_preview.json()["code"] == "invalid_batch_transition"
    assert again.status_code == 409 and again.json()["code"] == "invalid_batch_transition"


def test_rolling_back_requires_a_reason(pipeline: Pipeline) -> None:
    batch = _applied_update_and_create(pipeline)

    response = pipeline.login("ADMIN").post(f"{BASE}/{batch.id}/rollback", json={"reason": ""})

    assert response.status_code == 422
    assert get_batch(pipeline.client, batch.id)["status"] == "APPLIED"
