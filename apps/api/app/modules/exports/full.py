"""Full database export package (RF-044, design D4).

Produces a complete verifiable ZIP archive containing:
1. `csv/<entidad>.csv`: Every business entity, including soft-deleted rows with their audit stamps.
   Sensitive credentials (`password_hash`) are strictly excluded.
2. `matp_export.xlsx`: Multi-sheet workbook with one sheet per table.
3. `diccionario.csv`: Data dictionary (entity, column, type, Spanish description).
4. `manifest.json`: Manifest with timestamp, Alembic version, row counts and SHA-256 hashes.
"""

from __future__ import annotations

import csv
import datetime
import hashlib
import io
import json
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from sqlalchemy import inspect, text

from app.core.jobs import SessionFactory
from app.core.models_base import Base, new_uuid, utcnow
from app.core.storage import ObjectStorage
from app.modules.audit.context import AuditContext, set_audit_context
from app.modules.audit.models import AuditAction, AuditLog, AuditOrigin
from app.modules.exports.worker import EXPORT_TTL_HOURS, Heartbeat, _update_heartbeat
from app.modules.search.models import ExportJobModel, ExportJobStatus
from app.modules.users.models import AppUser

EXCLUDED_COLUMNS: dict[str, set[str]] = {
    "app_user": {"password_hash"},
}


def _format_export_cell(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.isoformat()
    if isinstance(val, uuid.UUID):
        return str(val)
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    return val


def export_full_database_job(
    factory: SessionFactory,
    job_id: uuid.UUID,
    storage: ObjectStorage | None = None,
) -> None:
    """Generate the full database verifiable export ZIP package."""
    heartbeat = Heartbeat(lambda: _update_heartbeat(factory, job_id))

    with tempfile.NamedTemporaryFile(suffix=".zip", delete=False) as tmp_zip:
        zip_path = Path(tmp_zip.name)

    try:
        manifest_files: list[dict[str, Any]] = []
        dictionary_rows: list[list[str]] = [["entidad", "columna", "tipo", "descripcion"]]
        total_rows_all_tables = 0

        # Entities to export (ordered)
        export_tables = [
            "piece",
            "piece_identifier",
            "collection",
            "vocabulary",
            "term",
            "location",
            "piece_movement",
            "loan",
            "loan_item",
            "media_asset",
            "audit_log",
            "app_user",
        ]

        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            wb = Workbook(write_only=True)

            with factory() as session:
                # Query Alembic revision head
                insp = inspect(session.connection())
                if insp.has_table("alembic_version"):
                    alembic_head = (
                        session.execute(text("SELECT version_num FROM alembic_version")).scalar()
                        or "0010_export_jobs"
                    )
                else:
                    alembic_head = "0010_export_jobs"

                for table_name in export_tables:
                    heartbeat.tick()
                    if table_name not in Base.metadata.tables:
                        continue

                    table = Base.metadata.tables[table_name]
                    excluded = EXCLUDED_COLUMNS.get(table_name, set())

                    # Filter out excluded columns (e.g. password_hash)
                    columns = [c.name for c in table.columns if c.name not in excluded]

                    # Document dictionary rows
                    for c in table.columns:
                        if c.name in excluded:
                            continue
                        col_type = str(c.type)
                        doc = c.comment or f"Columna {c.name} de {table_name}"
                        dictionary_rows.append([table_name, c.name, col_type, doc])

                    # Fetch rows (including soft deleted rows)
                    col_clause = ", ".join(f'"{c}"' for c in columns)
                    query = text(f"SELECT {col_clause} FROM {table_name}")
                    db_rows = session.execute(query).fetchall()
                    row_count = len(db_rows)
                    total_rows_all_tables += row_count

                    # 1. Write CSV
                    csv_buf = io.StringIO()
                    csv_writer = csv.writer(csv_buf)
                    csv_writer.writerow(columns)
                    for r in db_rows:
                        csv_writer.writerow([_format_export_cell(val) for val in r])

                    csv_bytes = csv_buf.getvalue().encode("utf-8-sig")
                    csv_path_in_zip = f"csv/{table_name}.csv"
                    zf.writestr(csv_path_in_zip, csv_bytes)

                    csv_sha256 = hashlib.sha256(csv_bytes).hexdigest()
                    manifest_files.append(
                        {
                            "path": csv_path_in_zip,
                            "row_count": row_count,
                            "sha256": csv_sha256,
                            "size_bytes": len(csv_bytes),
                        }
                    )

                    # 2. Add sheet to Excel workbook
                    ws = wb.create_sheet(title=table_name[:31])
                    ws.append(columns)
                    for r in db_rows:
                        ws.append([_format_export_cell(val) for val in r])

            # Save matp_export.xlsx
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_xlsx:
                xlsx_tmp_path = Path(tmp_xlsx.name)
            try:
                wb.save(xlsx_tmp_path)
                xlsx_bytes = xlsx_tmp_path.read_bytes()
                zf.writestr("matp_export.xlsx", xlsx_bytes)
                manifest_files.append(
                    {
                        "path": "matp_export.xlsx",
                        "row_count": total_rows_all_tables,
                        "sha256": hashlib.sha256(xlsx_bytes).hexdigest(),
                        "size_bytes": len(xlsx_bytes),
                    }
                )
            finally:
                if xlsx_tmp_path.exists():
                    xlsx_tmp_path.unlink()

            # 3. Write diccionario.csv
            dict_buf = io.StringIO()
            dict_writer = csv.writer(dict_buf)
            for d_row in dictionary_rows:
                dict_writer.writerow(d_row)
            dict_bytes = dict_buf.getvalue().encode("utf-8-sig")
            zf.writestr("diccionario.csv", dict_bytes)
            manifest_files.append(
                {
                    "path": "diccionario.csv",
                    "row_count": len(dictionary_rows) - 1,
                    "sha256": hashlib.sha256(dict_bytes).hexdigest(),
                    "size_bytes": len(dict_bytes),
                }
            )

            # 4. Write manifest.json
            manifest_data = {
                "format_version": "1.0",
                "generated_at": utcnow().isoformat(),
                "alembic_version": alembic_head,
                "total_entities": len(export_tables),
                "total_rows": total_rows_all_tables,
                "files": manifest_files,
            }
            manifest_bytes = json.dumps(manifest_data, indent=2).encode("utf-8")
            zf.writestr("manifest.json", manifest_bytes)

        # Upload zip to storage
        zip_bytes = zip_path.read_bytes()
        zip_sha256 = hashlib.sha256(zip_bytes).hexdigest()
        storage_key = f"exports/{job_id}/matp_full_export.zip"

        if storage is not None:
            storage.put_bytes(storage_key, zip_bytes, "application/zip")

        with factory() as session:
            job = session.get(ExportJobModel, job_id)
            if job:
                job.status = ExportJobStatus.READY
                job.storage_key = storage_key
                job.sha256 = zip_sha256
                job.row_count = total_rows_all_tables
                job.expires_at = utcnow() + datetime.timedelta(hours=EXPORT_TTL_HOURS)

                user = session.get(AppUser, job.requested_by_id)
                user_label = user.email if user else "admin"

                ctx = AuditContext(
                    origin=AuditOrigin.MANUAL if job.requested_by_id else AuditOrigin.SYSTEM,
                    user_id=job.requested_by_id,
                    actor_label=user_label or "admin",
                    action=AuditAction.EXPORT,
                )
                set_audit_context(session, ctx)

                audit_log = AuditLog(
                    change_set_id=new_uuid(),
                    entity_type="export_job",
                    entity_id=job.id,
                    field=None,
                    old_value=None,
                    new_value={
                        "kind": "full",
                        "format": "zip",
                        "row_count": total_rows_all_tables,
                        "sha256": zip_sha256,
                        "alembic_version": alembic_head,
                    },
                    action=AuditAction.EXPORT,
                    origin=AuditOrigin.MANUAL,
                    user_id=job.requested_by_id,
                    actor_label=user_label,
                    reason=f"Exportación completa de la base ({total_rows_all_tables} registros)",
                )
                session.add(audit_log)
                session.commit()

    finally:
        if zip_path.exists():
            zip_path.unlink()
