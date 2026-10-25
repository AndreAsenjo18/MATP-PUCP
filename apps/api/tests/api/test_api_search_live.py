"""Spec busqueda-reportes — búsqueda con el filtro único completo, miniatura de la foto principal
y selección de columnas al exportar (change busqueda-avanzada-y-exportacion, tareas 4.1 y 4.2)."""

import datetime
import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.models_base import utcnow
from app.models import Piece
from app.modules.media.models import MediaAsset
from app.modules.search.models import ExportJobKind, ExportJobModel, ExportJobStatus
from tests.api.conftest import LoginAs, Pipeline, SeededApi


def _ids(page: dict) -> list[str]:
    return [hit["piece"]["id"] for hit in page["items"]]


def test_search_without_text_lists_by_filters(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    page = client.get(
        "/api/v1/search",
        params={"tenure_regime": "Comodato", "has_inventory_code": "false", "limit": 100},
    )
    assert page.status_code == 200
    body = page.json()
    listed = client.get(
        "/api/v1/pieces",
        params={"tenure_regime": "Comodato", "has_inventory_code": "false", "limit": 100},
    ).json()
    assert body["total"] == listed["total"] > 0
    assert _ids(body) == [item["id"] for item in listed["items"]]
    assert all(hit.get("match_reason") is None for hit in body["items"])


def test_search_accepts_every_shared_filter(login_as: LoginAs) -> None:
    """Texto y filtros de `GET /pieces` combinados en AND: la misma selección que el listado."""
    client = login_as("CATALOGUER")
    sample = client.get("/api/v1/pieces", params={"limit": 100}).json()["items"]
    piece = next(item for item in sample if item["category"] and item["collection"])
    params = {
        "q": piece["denomination"].split()[0],
        "category_id": piece["category"]["id"],
        "collection_id": piece["collection"]["id"],
        "limit": 100,
    }
    found = client.get("/api/v1/search", params=params).json()
    listed = client.get("/api/v1/pieces", params=params).json()
    assert piece["id"] in _ids(found)
    assert set(_ids(found)) == {item["id"] for item in listed["items"]}
    assert all(hit["match_reason"] for hit in found["items"])


def test_search_alert_filters_split_the_catalog(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    total = client.get("/api/v1/search").json()["total"]
    with_alerts = client.get("/api/v1/search", params={"has_alerts": "true"}).json()["total"]
    without = client.get("/api/v1/search", params={"has_alerts": "false"}).json()["total"]
    assert with_alerts > 0
    assert with_alerts + without == total

    no_code = client.get(
        "/api/v1/search", params={"alert_code": "MISSING_INVENTORY_CODE", "limit": 100}
    ).json()
    assert no_code["total"] > 0
    for hit in no_code["items"]:
        assert hit["piece"]["tenure_regime"] == "Propiedad"
        assert hit["piece"]["code_i"] is None
    assert client.get("/api/v1/search", params={"alert_code": "NO_EXISTE"}).status_code == 422


def test_search_epoch_range_reports_uninterpreted_pieces(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    body = client.get(
        "/api/v1/search", params={"epoch_start_year": 1900, "epoch_end_year": 2000}
    ).json()
    assert isinstance(body["excluded_uninterpreted_period"], int)
    assert "excluded_uninterpreted_period" not in client.get("/api/v1/search").json()


def test_search_hit_carries_the_primary_photo_thumbnail(
    pipeline: Pipeline, monkeypatch: pytest.MonkeyPatch
) -> None:
    storage = pipeline.client.app.state.storage  # type: ignore[attr-defined]
    monkeypatch.setattr(storage, "presign_get", lambda key, ttl=300: f"https://s3.local/{key}")
    with pipeline.session() as session:
        with_media = select(MediaAsset.piece_id)
        piece = session.scalar(
            select(Piece).where(Piece.id.not_in(with_media)).order_by(Piece.denomination)
        )
        assert piece is not None
        piece_id, word = piece.id, piece.denomination.split()[0]
    with pipeline.session() as session:
        from app.modules.audit.context import AuditContext, set_audit_context
        from app.modules.audit.models import AuditOrigin

        set_audit_context(session, AuditContext(origin=AuditOrigin.SYSTEM, actor_label="tests"))
        primary = uuid.uuid4()
        for asset_id, is_primary, order in ((uuid.uuid4(), False, 0), (primary, True, 1)):
            session.add(
                MediaAsset(
                    id=asset_id,
                    piece_id=piece_id,
                    storage_key=f"pieces/{piece_id}/{asset_id}/original.jpg",
                    content_type="image/jpeg",
                    size_bytes=100,
                    content_sha256="a" * 64,
                    sort_order=order,
                    is_primary=is_primary,
                )
            )
        session.commit()

    page = pipeline.client.get("/api/v1/search", params={"q": word, "limit": 100}).json()
    hits = {hit["piece"]["id"]: hit for hit in page["items"]}
    assert hits[str(piece_id)]["thumbnail_url"] == (
        f"https://s3.local/pieces/{piece_id}/{primary}/thumb.webp"
    )
    others = [hit for pid, hit in hits.items() if pid != str(piece_id)]
    assert all(hit["thumbnail_url"] is None for hit in others if hit["piece"]["media_count"] == 0)


def _sheet(response) -> list[tuple]:
    workbook = load_workbook(io.BytesIO(response.content))
    return list(workbook["Resultados"].iter_rows(values_only=True))


def test_export_uses_the_selected_columns_in_order(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    response = client.post(
        "/api/v1/reports/export-excel",
        json={"filters": {"tenure_regime": "Comodato"}, "columns": ["denomination", "code_i"]},
    )
    assert response.status_code == 200
    rows = _sheet(response)
    assert rows[0] == ("Denominación", "Código I")
    listed = client.get("/api/v1/pieces", params={"tenure_regime": "Comodato"}).json()
    assert len(rows) - 1 == listed["total"]


def test_export_omits_sensitive_columns_the_role_cannot_see(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    response = client.post(
        "/api/v1/reports/export-excel",
        json={"filters": {}, "columns": ["denomination", "lender_name"]},
    )
    assert response.status_code == 200
    assert _sheet(response)[0] == ("Denominación",)
    assert "lender_name" in response.headers["x-omited-columns"]


def test_export_rejects_an_unknown_column(login_as: LoginAs) -> None:
    client = login_as("CATALOGUER")
    response = client.post(
        "/api/v1/reports/export-excel", json={"filters": {}, "columns": ["password_hash"]}
    )
    assert response.status_code == 422
    assert response.json()["code"] == "unknown_export_column"
    assert "denomination" in response.json()["details"]["allowed"]


def test_my_exports_show_kind_date_and_an_expired_file(seeded_api: SeededApi) -> None:
    client = seeded_api.login_as("CATALOGUER")
    user_id = uuid.UUID(client.get("/api/v1/auth/me").json()["id"])
    with seeded_api.session() as session:
        job = ExportJobModel(
            kind=ExportJobKind.RESULTS,
            requested_by_id=user_id,
            format="csv",
            params={"filters": {}},
            status=ExportJobStatus.READY,
            storage_key="exports/caducada/results.csv",
            row_count=12_000,
            expires_at=utcnow() - datetime.timedelta(hours=1),
        )
        session.add(job)
        session.commit()
        job_id = str(job.id)

    listed = {item["id"]: item for item in client.get("/api/v1/exports/jobs").json()}
    assert listed[job_id]["kind"] == "results"
    assert listed[job_id]["format"] == "csv"
    assert listed[job_id]["created_at"] is not None

    expired = client.get(f"/api/v1/exports/jobs/{job_id}/download-url")
    assert expired.status_code == 410
    assert expired.json()["code"] == "gone"
