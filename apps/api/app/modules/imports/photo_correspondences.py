"""Correspondences list of a photo-only batch (change importacion-fotos-a-catalogo, task 4.5; D7).

A CSV or Excel file with the columns ``archivo`` (file name or relative path) and ``codigo``, and
optionally ``tipo`` (identifier type), proposes the piece of each photo it names, with priority
over the reading of the name. Its lines are normalized and kept in
``import_batch.photo_reading.correspondences`` so a new reading respects them; the file itself is
not stored. A line whose file is not in the batch, whose code finds no piece, or whose file name
is repeated in several folders is reported with its reason without stopping the others. Manual
assignments are never changed (spec «Lista de correspondencias entre archivos y códigos»).
"""

import csv
import io
import uuid
from collections import defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ValidationFailed
from app.modules.identification.models import NormalizationRule
from app.modules.identification.normalization import fold_label
from app.modules.imports.catalog_photo_matching import (
    DUPLICATE_NAME,
    CatalogIndex,
    CatalogProposal,
    pieces_for_code,
)
from app.modules.imports.models import ImportBatch, ImportPhoto, ImportPhotoStatus
from app.modules.imports.models import PhotoMatchSource as Source
from app.modules.imports.models import PhotoMatchStatus as M

#: Header names, compared without case nor accents.
FILE_HEADERS = frozenset({"archivo", "foto", "file"})
CODE_HEADERS = frozenset({"codigo", "code"})
TYPE_HEADERS = frozenset({"tipo", "type"})
MAX_LINES = 20_000


@dataclass(frozen=True)
class Line:
    number: int  # line of the file, the header is 1
    file: str
    code: str
    type_code: str | None = None

    def as_json(self) -> dict[str, Any]:
        value: dict[str, Any] = {"line": self.number, "file": self.file, "code": self.code}
        if self.type_code:
            value["type"] = self.type_code
        return value

    @classmethod
    def from_json(cls, value: Mapping[str, Any]) -> "Line":
        return cls(int(value["line"]), value["file"], value["code"], value.get("type") or None)


def _rows(file_name: str, content: bytes) -> list[list[str]]:
    if file_name.casefold().endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises several kinds for a broken workbook
            raise ValidationFailed(
                "No se pudo leer el Excel de la lista.", code="invalid_correspondences"
            ) from exc
        sheet = book.worksheets[0]
        return [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def parse_correspondences(file_name: str, content: bytes) -> list[Line]:
    """The lines of the list, recognized by their headers (D7); 422 without them."""
    rows = [row for row in _rows(file_name, content) if any(row)]
    if not rows:
        raise ValidationFailed("La lista está vacía.", code="invalid_correspondences")
    headers = [fold_label(cell) for cell in rows[0]]

    def column(names: frozenset[str]) -> int | None:
        return next((i for i, header in enumerate(headers) if header in names), None)

    file_col, code_col, type_col = column(FILE_HEADERS), column(CODE_HEADERS), column(TYPE_HEADERS)
    if file_col is None or code_col is None:
        raise ValidationFailed(
            "La lista necesita las columnas «archivo» y «codigo» (y, si quiere, «tipo»).",
            code="invalid_correspondences",
            details={"headers": rows[0]},
        )
    if len(rows) - 1 > MAX_LINES:
        raise ValidationFailed(
            f"La lista supera las {MAX_LINES} líneas: divídala.", code="invalid_correspondences"
        )
    lines = []
    for number, row in enumerate(rows[1:], start=2):

        def cell(position: int | None, row: list[str] = row) -> str:
            return row[position] if position is not None and position < len(row) else ""

        file, code = cell(file_col).replace("\\", "/").strip("/"), cell(code_col)
        if file and code:
            lines.append(Line(number, file, code, cell(type_col).upper() or None))
    return lines


def merge_lines(previous: Sequence[Mapping[str, Any]], new: Sequence[Line]) -> list[Line]:
    """The lines kept in the batch: a new line for the same file replaces the old one."""
    by_file = {fold_label(item["file"]): Line.from_json(item) for item in previous}
    by_file.update({fold_label(line.file): line for line in new})
    return list(by_file.values())


@dataclass(frozen=True)
class LineResult:
    line: Line
    result: str  # applied | no_photo | no_piece | ambiguous
    reason: str | None = None
    photo_ids: tuple[uuid.UUID, ...] = ()


def resolve_lines(
    lines: Sequence[Line],
    photos: Sequence[ImportPhoto],
    index: CatalogIndex,
    rules: Mapping[str, NormalizationRule],
) -> tuple[dict[uuid.UUID, CatalogProposal], list[LineResult]]:
    """The proposal of every photo a line names and the result of each line (D7)."""
    by_path: dict[str, list[ImportPhoto]] = defaultdict(list)
    by_name: dict[str, list[ImportPhoto]] = defaultdict(list)
    for photo in photos:
        by_path[fold_label(photo.relative_path)].append(photo)
        by_name[fold_label(photo.original_filename)].append(photo)
    preset: dict[uuid.UUID, CatalogProposal] = {}
    results: list[LineResult] = []
    for line in lines:
        key = fold_label(line.file)
        named = by_path.get(key) or ([] if "/" in line.file else by_name.get(key, []))
        if not named:
            results.append(LineResult(line, "no_photo", "La foto no está en el lote."))
            continue
        ids = tuple(photo.id for photo in named)
        if line.type_code is not None and line.type_code not in rules:
            results.append(LineResult(line, "no_piece", f"«{line.type_code}» no es un tipo.", ids))
            continue
        pieces = pieces_for_code(line.code, index, type_code=line.type_code, rules=rules)
        if not pieces:
            results.append(LineResult(line, "no_piece", "Ninguna pieza tiene ese código.", ids))
            continue
        candidates = tuple(sorted(pieces, key=str))
        if len(named) > 1:
            for photo in named:
                preset[photo.id] = CatalogProposal(
                    photo.id,
                    M.AMBIGUOUS,
                    Source.LIST,
                    candidate_piece_ids=candidates,
                    reason=DUPLICATE_NAME,
                )
            results.append(
                LineResult(
                    line, "ambiguous", "Nombre repetido en el lote: indique la ruta completa.", ids
                )
            )
            continue
        if len(pieces) > 1:
            preset[ids[0]] = CatalogProposal(
                ids[0], M.AMBIGUOUS, Source.LIST, candidate_piece_ids=candidates
            )
            results.append(LineResult(line, "ambiguous", "El código está en varias piezas.", ids))
            continue
        preset[ids[0]] = CatalogProposal(ids[0], M.MATCHED, Source.LIST, candidates[0])
        results.append(LineResult(line, "applied", None, ids))
    return preset, results


def batch_preset(
    session: Session,
    batch: ImportBatch,
    index: CatalogIndex,
    rules: Mapping[str, NormalizationRule],
) -> tuple[dict[uuid.UUID, CatalogProposal], list[LineResult]]:
    """The proposals of the lines kept in the batch, for every run of the matching."""
    stored = (batch.photo_reading or {}).get("correspondences", [])
    lines = [Line.from_json(item) for item in stored]
    if not lines:
        return {}, []
    photos = list(
        session.scalars(
            select(ImportPhoto).where(
                ImportPhoto.batch_id == batch.id,
                ImportPhoto.status == ImportPhotoStatus.VERIFIED,
            )
        )
    )
    return resolve_lines(lines, photos, index, rules)
