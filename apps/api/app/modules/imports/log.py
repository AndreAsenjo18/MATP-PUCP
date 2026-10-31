"""Load log of an import batch (spec importacion-datos: RF-028).

The log gathers the file, sheet, mapping, uploader, approver, stage timestamps, counts per
classification and decision, skipped empty rows and, for every row that is not applied, its
source row number and reason. Rows a person resolved are listed too, with the resolution, and
the decisions of a person on the photos of the batch (who, when, what and on how many photos;
change importacion-fotos-a-catalogo, task 4.1), read from the audit log by the batch.
It is offered as JSON and as a spreadsheet with one sheet for the summary and one per list.
"""

import io
import uuid
from typing import Any

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.modules.audit.models import AuditLog
from app.modules.imports.decisions import RowAction, row_action
from app.modules.imports.models import ImportBatch, ImportRow
from app.modules.imports.models import ImportBatchStatus as S
from app.modules.imports.state import STATUS_LABELS
from app.modules.users.models import AppUser

APPLIED_STATUSES = frozenset({S.APPLIED, S.REVERTED})
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _user_name(session: Session, user_id: uuid.UUID | None) -> str | None:
    if user_id is None:
        return None
    user = session.get(AppUser, user_id)
    return user.full_name if user is not None else None


def photo_decisions(session: Session, batch_id: uuid.UUID) -> list[dict[str, Any]]:
    """Decisions of a person on the photos of the batch, one per change set (task 4.1)."""
    rows = session.execute(
        select(
            AuditLog.change_set_id,
            func.min(AuditLog.occurred_at),
            AuditLog.user_id,
            AuditLog.reason,
            func.count(func.distinct(AuditLog.entity_id)),
        )
        .where(
            AuditLog.entity_type == "import_photo",
            AuditLog.origin_ref == str(batch_id),
            AuditLog.user_id.is_not(None),
        )
        .group_by(AuditLog.change_set_id, AuditLog.user_id, AuditLog.reason)
        .order_by(func.min(AuditLog.occurred_at))
    ).all()
    return [
        {
            "at": occurred_at.isoformat(),
            "user": _user_name(session, user_id),
            "operation": reason,
            "photos": int(count),
        }
        for _, occurred_at, user_id, reason, count in rows
    ]


def build_log(session: Session, batch: ImportBatch) -> dict[str, Any]:
    raw = batch.counts or {}
    rows = session.scalars(
        select(ImportRow)
        .where(ImportRow.batch_id == batch.id)
        .order_by(ImportRow.source_row_number)
    ).all()
    entries: list[dict[str, Any]] = []
    skipped = 0
    for row in rows:
        action, reason = row_action(row)
        skipped += int(action == RowAction.SKIP)
        if action != RowAction.SKIP and not row.resolution:
            continue
        entries.append(
            {
                "source_row_number": row.source_row_number,
                "classification": row.classification,
                "decision": row.decision,
                "action": action.value,
                "reason": reason,
                "resolution": row.resolution,
                "target_piece_id": row.target_piece_id,
            }
        )
    applied = batch.status in APPLIED_STATUSES
    has_rows = bool(rows)
    return {
        "batch": {
            "id": str(batch.id),
            "status": batch.status.value,
            "status_label": STATUS_LABELS[batch.status],
            "status_reason": batch.status_reason,
            "source_name": batch.source_name,
            "file_name": batch.file_name,
            "file_sha256": batch.file_sha256,
            "sheet_name": batch.sheet_name,
            "header_row": batch.header_row,
            "template_id": str(batch.template_id) if batch.template_id else None,
            "mapping": batch.mapping,
            "uploaded_by": _user_name(session, batch.uploaded_by_id),
            "approved_by": _user_name(session, batch.approved_by_id),
            "approved_at": batch.approved_at.isoformat() if batch.approved_at else None,
            "stage_timestamps": batch.stage_timestamps or {},
            "counts": {
                key: value
                for key, value in raw.items()
                if isinstance(value, int) and not isinstance(value, bool)
            }
            if has_rows
            else {},
            "skipped_empty_rows": raw.get("skipped_empty_rows") if has_rows else None,
        },
        "applied_rows": (
            int(raw.get("APPLIED_CREATED", 0)) + int(raw.get("APPLIED_UPDATED", 0))
            if applied
            else None
        ),
        "rejected_rows": skipped if has_rows else None,
        "entries": entries,
        "photo_decisions": photo_decisions(session, batch.id),
    }


_SUMMARY_LABELS = (
    ("file_name", "Archivo"),
    ("source_name", "Fuente"),
    ("sheet_name", "Hoja"),
    ("header_row", "Fila de cabeceras"),
    ("status_label", "Estado"),
    ("status_reason", "Motivo del estado"),
    ("uploaded_by", "Cargado por"),
    ("approved_by", "Aprobado por"),
    ("approved_at", "Fecha de aprobación"),
    ("skipped_empty_rows", "Filas vacías omitidas"),
)
_ROW_HEADERS = ("Fila", "Clasificación", "Decisión", "Acción", "Motivo", "Pieza")


def _entry_values(entry: dict[str, Any]) -> list[Any]:
    return [
        entry["source_row_number"],
        entry["classification"].value if entry["classification"] else None,
        entry["decision"].value,
        entry["action"],
        entry["reason"],
        str(entry["target_piece_id"]) if entry["target_piece_id"] else None,
    ]


def log_workbook(log: dict[str, Any]) -> bytes:
    """Spreadsheet of the log: summary, rows not applied and rows resolved by a person."""
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Resumen"
    batch = log["batch"]
    for key, label in _SUMMARY_LABELS:
        summary.append([label, batch.get(key)])
    for stage, moment in (batch.get("stage_timestamps") or {}).items():
        summary.append([f"Etapa {stage}", moment])
    for key, value in (batch.get("counts") or {}).items():
        summary.append([f"Conteo {key}", value])
    summary.append(["Filas aplicadas", log["applied_rows"]])
    summary.append(["Filas no aplicadas", log["rejected_rows"]])

    skipped = workbook.create_sheet("Filas no aplicadas")
    skipped.append(list(_ROW_HEADERS))
    resolved = workbook.create_sheet("Resoluciones")
    resolved.append([*_ROW_HEADERS, "Resolución"])
    for entry in log["entries"]:
        if entry["action"] == RowAction.SKIP.value:
            skipped.append(_entry_values(entry))
        if entry["resolution"]:
            resolution = entry["resolution"]
            detail = (
                "Nueva pieza"
                if resolution.get("as_new")
                else f"Pieza destino {resolution['target_piece_id']}"
                if resolution.get("target_piece_id")
                else None
            )
            resolved.append([*_entry_values(entry), detail])
    if log.get("photo_decisions"):
        decisions = workbook.create_sheet("Decisiones sobre fotos")
        decisions.append(["Fecha", "Usuario", "Decisión", "Fotos"])
        for decision in log["photo_decisions"]:
            decisions.append(
                [decision["at"], decision["user"], decision["operation"], decision["photos"]]
            )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
