"""Spec plataforma — Datos de demostración sintéticos reproducibles (RNF-008, RNF-014)."""

from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine, create_engine, event, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.db import make_session_factory
from app.models import (
    AiSuggestion,
    AppUser,
    AuditLog,
    Base,
    DuplicateCandidate,
    ImportMappingTemplate,
    MediaAsset,
    Piece,
    PieceIdentifier,
)
from app.modules.audit.models import AuditOrigin
from app.modules.audit.tracking import INCLUDE_DELETED
from app.modules.catalog.enums import TenureRegime
from app.modules.imports.mapping import (
    MappingContext,
    MappingSpecV1,
    read_mapping,
    validate_mapping,
)
from app.modules.imports.templates import header_signature
from app.seed.demo_batch import HEADERS as DEMO_BATCH_HEADERS
from app.seed.fixtures import HEADERS, write_synthetic_workbook
from app.seed.import_templates import PREDEFINED_TEMPLATES
from app.seed.synthetic import SeedOptions, SeedRefused, run_seed

ALL = {INCLUDE_DELETED: True}


def _new_session() -> tuple[Engine, Session]:
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    event.listen(engine, "connect", lambda conn, _r: conn.execute("PRAGMA foreign_keys=ON"))
    factory = make_session_factory(engine)
    Base.metadata.create_all(engine)
    return engine, factory()


@pytest.fixture(scope="module")
def seeded() -> Iterator[tuple[Session, Any]]:
    engine, session = _new_session()
    summary = run_seed(session, SeedOptions(upload_media=False, admin_password="demo-password-1"))
    session.commit()
    yield session, summary
    session.close()
    engine.dispose()


def _count(session: Session, statement: Any) -> int:
    return session.scalar(statement.execution_options(**ALL)) or 0


def test_demo_catalogue_size_and_distribution(seeded: tuple[Session, Any]) -> None:
    session, summary = seeded
    pieces = _count(session, select(func.count()).select_from(Piece))
    assert pieces == summary.pieces == 300
    with_inventory = select(PieceIdentifier.piece_id).where(
        PieceIdentifier.identifier_type_code == "I", PieceIdentifier.is_current
    )
    without_i = _count(
        session, select(func.count()).select_from(Piece).where(Piece.id.not_in(with_inventory))
    )
    assert 0.4 <= without_i / pieces <= 0.6
    assert summary.loan_for_use_pieces > 0
    assert summary.loose_pieces > 0
    assert summary.collections == 6
    assert 0 < summary.located_pieces < pieces  # partial locations (RF-019)


def test_no_loan_for_use_or_temporary_loan_piece_has_inventory_code(
    seeded: tuple[Session, Any],
) -> None:
    session, _ = seeded
    offending = _count(
        session,
        select(func.count())
        .select_from(PieceIdentifier)
        .join(Piece, Piece.id == PieceIdentifier.piece_id)
        .where(
            PieceIdentifier.identifier_type_code == "I",
            Piece.tenure_regime != TenureRegime.OWNED,
        ),
    )
    assert offending == 0


def test_dirty_codes_are_kept_as_originals(seeded: tuple[Session, Any]) -> None:
    session, _ = seeded
    identifiers = session.scalars(select(PieceIdentifier).execution_options(**ALL)).all()
    dirty = [i for i in identifiers if i.original_value != i.normalized_value]
    assert any("." in i.original_value for i in dirty if i.identifier_type_code == "COLECCION")
    assert any(i.normalization_status.value == "UNPARSEABLE" for i in identifiers)
    assert any(not i.is_current for i in identifiers)  # history (INC 4 -> 6 digits, correction)


def test_duplicates_photos_and_history(seeded: tuple[Session, Any]) -> None:
    session, summary = seeded
    assert _count(session, select(func.count()).select_from(DuplicateCandidate)) >= 10
    per_piece = session.execute(
        select(MediaAsset.piece_id, func.count()).group_by(MediaAsset.piece_id)
    ).all()
    assert per_piece and max(count for _, count in per_piece) > 1
    assert summary.uploaded_media == 0
    assert summary.soft_deleted_pieces == 2


OLD_FIELD_NAMES = {"title", "period_text", "period_from", "period_to", "period_type"}


def _json_keys_and_strings(value: Any) -> Iterator[str]:
    if isinstance(value, dict):
        for key, item in value.items():
            yield key
            yield from _json_keys_and_strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from _json_keys_and_strings(item)
    elif isinstance(value, str):
        yield value


def test_seed_json_uses_contract_field_names(seeded: tuple[Session, Any]) -> None:
    """Task 2.4: same volumes as before and no pre-contract field names in seed JSON."""
    session, summary = seeded
    assert summary.pieces == 300
    assert _count(session, select(func.count()).select_from(DuplicateCandidate)) == 10
    assert 0.4 <= summary.pieces_without_inventory_code / summary.pieces <= 0.6
    documents: list[Any] = []
    documents += session.scalars(select(DuplicateCandidate.matched_fields)).all()
    documents += session.scalars(select(ImportMappingTemplate.mapping)).all()
    documents += session.scalars(select(AiSuggestion.input_data)).all()
    assert len(documents) > 10
    found = {text for text in _json_keys_and_strings(documents) if text in OLD_FIELD_NAMES}
    assert found == set()


def test_every_seeded_change_is_audited_as_system_or_synthetic_user(
    seeded: tuple[Session, Any],
) -> None:
    session, _ = seeded
    origins = set(session.scalars(select(AuditLog.origin).distinct()))
    assert AuditOrigin.SYSTEM in origins
    assert _count(session, select(func.count()).select_from(AuditLog)) > 1000


def test_users_are_synthetic(seeded: tuple[Session, Any]) -> None:
    session, _ = seeded
    users = session.scalars(select(AppUser)).all()
    assert users and all(user.is_synthetic for user in users)
    assert all(user.email.endswith("@matp.local") for user in users)
    admin = next(user for user in users if user.email == "admin@matp.local")
    assert admin.password_hash and "demo-password-1" not in admin.password_hash


def test_seed_refuses_non_empty_catalogue(seeded: tuple[Session, Any]) -> None:
    session, _ = seeded
    with pytest.raises(SeedRefused) as excinfo:
        run_seed(session, SeedOptions(upload_media=False))
    assert "npm run" in str(excinfo.value)


def _catalogue_fingerprint(options: SeedOptions) -> list[tuple[str, str, tuple[str, ...]]]:
    engine, session = _new_session()
    try:
        run_seed(session, options)
        session.commit()
        pieces = session.scalars(select(Piece).execution_options(**ALL)).all()
        rows = []
        for piece in pieces:
            codes = tuple(sorted(i.original_value for i in piece.identifiers))
            rows.append((piece.denomination, piece.tenure_regime.value, codes))
        return sorted(rows)
    finally:
        session.close()
        engine.dispose()


def test_seed_is_reproducible_with_same_random_seed() -> None:
    options = SeedOptions(piece_count=60, upload_media=False, duplicate_pairs=3, sets=1)
    assert _catalogue_fingerprint(options) == _catalogue_fingerprint(options)


class FakeStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}
        self.bucket_checked = False

    def ensure_bucket(self) -> bool:
        self.bucket_checked = True
        return True

    def put_bytes(self, key: str, data: bytes, content_type: str) -> None:
        assert content_type == "image/jpeg"
        self.objects[key] = data


def test_photos_are_uploaded_to_object_storage_not_database() -> None:
    engine, session = _new_session()
    storage = FakeStorage()
    try:
        summary = run_seed(
            session,
            SeedOptions(piece_count=40, duplicate_pairs=2, sets=1),
            storage,  # type: ignore[arg-type]
        )
        session.commit()
        keys = set(session.scalars(select(MediaAsset.storage_key)))
    finally:
        session.close()
        engine.dispose()
    assert storage.bucket_checked
    assert summary.uploaded_media == summary.media_assets == len(storage.objects) > 0
    assert keys == set(storage.objects)
    assert all(data[:2] == b"\xff\xd8" for data in storage.objects.values())  # JPEG


def test_synthetic_excel_fixture(tmp_path: Path) -> None:
    path = write_synthetic_workbook(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["SABANA"]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row >= 41
    values = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "I 2362 / RA 28" in values and "S/N" in values
    assert len(sheet._images) == 3  # openpyxl exposes images only privately


FIXTURES = Path(__file__).resolve().parents[3] / "data" / "fixtures"
SABANA_FIXTURE = FIXTURES / "sabana_sintetica_v1.xlsx"


def test_seed_template_is_v1_for_the_synthetic_sabana(seeded: tuple[Session, Any]) -> None:
    """Task 3.4 of plantillas-mapeo-y-normalizacion: the seed template is v1, marked
    [SUPUESTO], with the signature of the headers of ``data/fixtures/sabana_sintetica_v1.xlsx``
    computed as the API does, and its mapping is valid for those headers."""
    session, _summary = seeded
    template = session.scalars(
        select(ImportMappingTemplate).where(
            ImportMappingTemplate.name == "Plantilla sábana sintética v1"
        )
    ).one()
    sheet = load_workbook(SABANA_FIXTURE, read_only=True)["SABANA"]
    headers = [str(cell.value) for cell in next(sheet.iter_rows(max_row=1))]

    assert template.mapping["version"] == 1
    assert "[SUPUESTO]" in (template.description or "")
    assert headers == HEADERS
    assert template.header_signature == header_signature(headers)
    spec = read_mapping(template.mapping)
    assert isinstance(spec, MappingSpecV1)
    assert [column.source for column in spec.columns] == headers
    assert validate_mapping(spec, headers, MappingContext(session)) == []


def test_seed_has_a_template_that_recognizes_the_demo_batch(seeded: tuple[Session, Any]) -> None:
    """Change rediseno-ux-importacion: the demo batch (``data/fixtures/lote_demo``) is recognized
    by its headers, and its "RÉGIMEN" column reaches the validation as the tenure regime."""
    session, _summary = seeded
    template = session.scalars(
        select(ImportMappingTemplate).where(ImportMappingTemplate.name == "Lote de demostración v1")
    ).one()

    spec = read_mapping(template.mapping)

    assert "[SUPUESTO]" in (template.description or "")
    assert template.header_signature == header_signature(DEMO_BATCH_HEADERS)
    assert isinstance(spec, MappingSpecV1)
    assert [column.source for column in spec.columns] == DEMO_BATCH_HEADERS
    assert validate_mapping(spec, DEMO_BATCH_HEADERS, MappingContext(session)) == []
    assert any(
        column.source == "RÉGIMEN" and column.target.model_dump().get("name") == "tenure_regime"
        for column in spec.columns
    )


def test_seed_has_one_predefined_template_per_source_the_museum_names(
    seeded: tuple[Session, Any],
) -> None:
    """Change rediseno-ux-administracion: Administration shows more formats than the sábana ones,
    each an illustrative [SUPUESTO] v1 template whose mapping is valid for its own headers."""
    session, _summary = seeded
    stored = {
        template.name: template for template in session.scalars(select(ImportMappingTemplate)).all()
    }

    assert len(PREDEFINED_TEMPLATES) == 5
    assert len({template.name for template in PREDEFINED_TEMPLATES}) == 5
    assert len({template.source_name for template in PREDEFINED_TEMPLATES}) == 5
    for predefined in PREDEFINED_TEMPLATES:
        template = stored[predefined.name]
        spec = read_mapping(template.mapping)

        assert "[SUPUESTO]" in (template.description or "")
        assert template.source_name == predefined.source_name
        assert template.header_signature == header_signature(predefined.headers)
        assert isinstance(spec, MappingSpecV1)
        assert [column.source for column in spec.columns] == predefined.headers
        assert validate_mapping(spec, predefined.headers, MappingContext(session)) == []
        assert any(
            column.target.kind == "field" and column.target.name == "denomination"
            for column in spec.columns
        )


def test_the_predefined_templates_do_not_look_alike_so_detection_is_unambiguous() -> None:
    """A file with the headers of one of the five formats must not reach the suggestion threshold
    (0.6) of any other format: the wizard applies the best match without asking. (The sábana and
    demo batch formats already share most columns; the wizard picks the best of those.)"""
    from app.modules.imports.templates import SUGGEST_THRESHOLD, header_key, similarity

    all_headers = [template.headers for template in PREDEFINED_TEMPLATES] + [
        HEADERS,
        DEMO_BATCH_HEADERS,
    ]
    keys = [{header_key(header) for header in headers} for headers in all_headers]
    for index in range(len(PREDEFINED_TEMPLATES)):
        for other_index, other in enumerate(keys):
            if index != other_index:
                assert similarity(keys[index], other) < SUGGEST_THRESHOLD, (
                    all_headers[index],
                    all_headers[other_index],
                )
