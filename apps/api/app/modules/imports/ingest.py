"""Spreadsheet ingestion for the import pipeline (spec importacion-datos: RF-021, RF-028; D2).

Pure functions, no database access: they turn the bytes of an uploaded ``.xlsx`` or ``.csv`` into
headers and raw rows. The caller stores the original file untouched and persists the rows.

- ``.xlsx``: ``openpyxl`` in ``read_only`` mode with ``data_only=True``, so formulas yield the
  value calculated and stored by the spreadsheet program. The sheet is chosen by the user
  (default: first non-empty sheet). Merged cells propagate the value of their top-left cell to
  every covered cell and the affected columns are listed in ``raw_data["__merged"]``.
- ``.csv``: encoding UTF-8 (with or without BOM), then Windows-1252, then Latin-1; separator
  ``;``, ``,``, tab or ``|`` detected with ``csv.Sniffer``.
- The header row is given by the user (1-based) or detected [SUPUESTO]: the first row, among the
  first ``HEADER_SCAN_ROWS``, with at least two text cells and no numbers.
- Completely empty rows between data rows are skipped and counted; trailing empty rows are not.
- Formats other than ``.xlsx``/``.csv`` are rejected [SUPUESTO A1].
"""

import csv
import io
import posixpath
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import PurePath
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import ValidationFailed

SUPPORTED_EXTENSIONS = (".xlsx", ".csv")
HEADER_SCAN_ROWS = 30
MERGED_KEY = "__merged"
CSV_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
CSV_DELIMITERS = ";,\t|"

_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_NS_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


class IngestionFailed(ValidationFailed):
    """The file cannot be ingested; the batch goes to ``FAILED_INGESTION`` with this reason."""

    code = "ingestion_failed"


@dataclass
class IngestedRow:
    row_number: int  # as the user sees it in the spreadsheet (1-based) or CSV record number
    raw_data: dict[str, Any]


@dataclass
class IngestResult:
    file_format: str  # "xlsx" | "csv"
    sheet_names: list[str]
    sheet_name: str | None
    header_row: int
    headers: list[str]
    rows: list[IngestedRow] = field(default_factory=list)
    skipped_empty_rows: int = 0
    encoding: str | None = None
    delimiter: str | None = None

    @property
    def row_count(self) -> int:
        return len(self.rows)


def file_format(filename: str) -> str:
    """Return ``xlsx`` or ``csv`` from the file name, or fail with ``unsupported_format``."""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise IngestionFailed(
            f"Formato no admitido ({suffix or 'sin extensión'}). Solo se aceptan archivos "
            ".xlsx y .csv; convierta el archivo antes de subirlo.",
            code="unsupported_format",
            details={"extension": suffix, "supported": list(SUPPORTED_EXTENSIONS)},
        )
    return suffix[1:]


def ingest(
    content: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    header_row: int | None = None,
    max_rows: int | None = None,
) -> IngestResult:
    """Read headers and raw rows from an ``.xlsx`` or ``.csv`` file.

    Raises ``IngestionFailed`` (``unsupported_format``, ``unreadable_file``, ``sheet_not_found``,
    ``empty_file``, ``header_row_not_found`` or ``row_limit_exceeded``).
    """
    if header_row is not None and header_row < 1:
        raise IngestionFailed(
            "La fila de cabeceras debe ser 1 o mayor.",
            code="header_row_not_found",
            details={"header_row": header_row},
        )
    fmt = file_format(filename)
    if fmt == "xlsx":
        return _ingest_xlsx(
            content, sheet_name=sheet_name, header_row=header_row, max_rows=max_rows
        )
    return _ingest_csv(content, header_row=header_row, max_rows=max_rows)


def list_sheets(content: bytes) -> list[str]:
    """Sheet names of an ``.xlsx`` file, in workbook order."""
    workbook = _open_workbook(content)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


# --- xlsx -------------------------------------------------------------------------------------
def _open_workbook(content: bytes):
    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise IngestionFailed(
            "No se pudo leer el archivo Excel; puede estar dañado o protegido con contraseña.",
            code="unreadable_file",
        ) from exc


def _ingest_xlsx(
    content: bytes, *, sheet_name: str | None, header_row: int | None, max_rows: int | None
) -> IngestResult:
    workbook = _open_workbook(content)
    try:
        sheet_names = list(workbook.sheetnames)
        if sheet_name is not None and sheet_name not in sheet_names:
            raise IngestionFailed(
                f"El archivo no tiene una hoja llamada «{sheet_name}».",
                code="sheet_not_found",
                details={"sheet_name": sheet_name, "sheet_names": sheet_names},
            )
        candidates = [sheet_name] if sheet_name is not None else sheet_names
        for name in candidates:
            worksheet = workbook[name]
            worksheet.reset_dimensions()  # do not trust the stored dimension (may truncate rows)
            rows = (
                tuple(_cell_value(v) for v in row) for row in worksheet.iter_rows(values_only=True)
            )
            merged = _merged_cell_map(content, name)
            result = _build_result(
                _propagate_merged(rows, merged),
                header_row=header_row,
                max_rows=max_rows,
                fail_if_empty=sheet_name is not None,
            )
            if result is not None:
                result.file_format = "xlsx"
                result.sheet_names = sheet_names
                result.sheet_name = name
                return result
        raise IngestionFailed(
            "El archivo no tiene datos en ninguna hoja.",
            code="empty_file",
            details={"sheet_names": sheet_names},
        )
    finally:
        workbook.close()


def _cell_value(value: Any) -> Any:
    """JSON-friendly raw value; the original file keeps the exact cell."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value.strip() else None
    if isinstance(value, bool | int | float):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def _merged_cell_map(content: bytes, sheet_name: str) -> dict[tuple[int, int], tuple[int, int]]:
    """Map every covered cell ``(row, col)`` to the top-left cell of its merged range.

    ``read_only`` worksheets do not expose merged ranges, so ``<mergeCells>`` is read from the
    sheet XML of the package.
    """
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        path = _sheet_path(package, sheet_name)
        if path is None:
            return {}
        covered: dict[tuple[int, int], tuple[int, int]] = {}
        with package.open(path) as stream:
            for _, element in ElementTree.iterparse(stream):
                if element.tag == f"{_NS_MAIN}mergeCell":
                    min_col, min_row, max_col, max_row = range_boundaries(element.get("ref", ""))
                    for row in range(min_row, max_row + 1):
                        for col in range(min_col, max_col + 1):
                            if (row, col) != (min_row, min_col):
                                covered[(row, col)] = (min_row, min_col)
                elif element.tag == f"{_NS_MAIN}row":
                    element.clear()  # keep memory flat on large sheets
        return covered


def _sheet_path(package: zipfile.ZipFile, sheet_name: str) -> str | None:
    workbook = ElementTree.fromstring(package.read("xl/workbook.xml"))
    rel_id = None
    for sheet in workbook.iter(f"{_NS_MAIN}sheet"):
        if sheet.get("name") == sheet_name:
            rel_id = sheet.get(f"{_NS_REL}id")
    if rel_id is None:
        return None
    rels = ElementTree.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    for rel in rels.iter(f"{_NS_PKG_REL}Relationship"):
        if rel.get("Id") == rel_id:
            target = rel.get("Target", "")
            return (
                target.lstrip("/") if target.startswith("/") else posixpath.normpath(f"xl/{target}")
            )
    return None


def _propagate_merged(
    rows: Iterable[tuple[Any, ...]], merged: dict[tuple[int, int], tuple[int, int]]
) -> Iterator[tuple[tuple[Any, ...], frozenset[int]]]:
    """Yield ``(values, merged_columns)``; covered cells take the value of their anchor."""
    anchor_cols: dict[int, set[int]] = {}
    for anchor_row, anchor_col in merged.values():
        anchor_cols.setdefault(anchor_row, set()).add(anchor_col)
    covered_by_row: dict[int, list[tuple[int, tuple[int, int]]]] = {}
    for (row, col), anchor in merged.items():
        covered_by_row.setdefault(row, []).append((col, anchor))
    anchor_values: dict[tuple[int, int], Any] = {}
    for row_number, values in enumerate(rows, start=1):
        for col in anchor_cols.get(row_number, ()):
            if col <= len(values):
                anchor_values[(row_number, col)] = values[col - 1]
        covered = covered_by_row.get(row_number)
        if not covered:
            yield values, frozenset()
            continue
        filled = list(values)
        merged_cols: set[int] = set()
        for col, anchor in covered:
            value = anchor_values.get(anchor)
            if value is None:
                continue
            while len(filled) < col:
                filled.append(None)
            filled[col - 1] = value
            merged_cols.add(col)
        yield tuple(filled), frozenset(merged_cols)


# --- csv --------------------------------------------------------------------------------------
def _ingest_csv(content: bytes, *, header_row: int | None, max_rows: int | None) -> IngestResult:
    text, encoding = _decode(content)
    if not text.strip():
        raise IngestionFailed("El archivo está vacío.", code="empty_file")
    delimiter = _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = ((tuple(v if v.strip() else None for v in record), frozenset()) for record in reader)
    result = _build_result(rows, header_row=header_row, max_rows=max_rows, fail_if_empty=True)
    assert result is not None
    result.file_format = "csv"
    result.encoding = encoding
    result.delimiter = delimiter
    return result


def _decode(content: bytes) -> tuple[str, str]:
    for encoding in CSV_ENCODINGS:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise AssertionError("latin-1 decodes any byte sequence")  # pragma: no cover


def _detect_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS).delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample else ""
        return max(CSV_DELIMITERS, key=first_line.count) if first_line else ","


# --- common -----------------------------------------------------------------------------------
def _is_empty(values: tuple[Any, ...]) -> bool:
    return all(value is None for value in values)


def _own_values(values: tuple[Any, ...], merged_cols: frozenset[int]) -> list[Any]:
    """Non-empty values of the row itself, without those propagated from a merged range."""
    return [v for i, v in enumerate(values, start=1) if v is not None and i not in merged_cols]


def _detect_header_row(scanned: list[tuple[tuple[Any, ...], frozenset[int]]]) -> int | None:
    for index, (values, merged_cols) in enumerate(scanned, start=1):
        filled = _own_values(values, merged_cols)
        if len(filled) >= 2 and all(isinstance(value, str) for value in filled):
            return index
    for index, (values, _) in enumerate(scanned, start=1):
        if not _is_empty(values):
            return index
    return None


def _header_names(values: tuple[Any, ...]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    last = max((i for i, v in enumerate(values) if v is not None), default=-1)
    for index, value in enumerate(values[: last + 1], start=1):
        name = str(value).strip() if value is not None else ""
        names.append(_unique(name or _generated_name(index), seen))
    return names


def _generated_name(column: int) -> str:
    return f"Columna {get_column_letter(column)}"


def _unique(name: str, seen: dict[str, int]) -> str:
    count = seen.get(name, 0) + 1
    seen[name] = count
    if count == 1:
        return name
    return _unique(f"{name} ({count})", seen)


def _build_result(
    rows: Iterable[tuple[tuple[Any, ...], frozenset[int]]],
    *,
    header_row: int | None,
    max_rows: int | None,
    fail_if_empty: bool,
) -> IngestResult | None:
    """Turn ``(values, merged_columns)`` rows into headers and data rows.

    Returns ``None`` when no header row is found and ``fail_if_empty`` is false (the caller
    tries the next sheet).
    """
    iterator = iter(rows)
    scanned: list[tuple[tuple[Any, ...], frozenset[int]]] = []
    limit = header_row if header_row is not None else HEADER_SCAN_ROWS
    for item in iterator:
        scanned.append(item)
        if len(scanned) >= limit:
            break

    if header_row is None:
        found = _detect_header_row(scanned)
    else:
        found = (
            header_row
            if len(scanned) >= header_row and not _is_empty(scanned[header_row - 1][0])
            else None
        )
    if found is None:
        if not fail_if_empty:
            return None
        if header_row is None:
            raise IngestionFailed("El archivo no tiene datos.", code="empty_file")
        raise IngestionFailed(
            f"La fila {header_row} no existe o está vacía; indique la fila de cabeceras.",
            code="header_row_not_found",
            details={"header_row": header_row},
        )

    headers = _header_names(scanned[found - 1][0])
    seen = {name: 1 for name in headers}
    result = IngestResult(
        file_format="", sheet_names=[], sheet_name=None, header_row=found, headers=headers
    )
    pending_empty = 0

    def data_rows() -> Iterator[tuple[int, tuple[tuple[Any, ...], frozenset[int]]]]:
        yield from enumerate(scanned[found:], start=found + 1)
        yield from enumerate(iterator, start=len(scanned) + 1)

    for row_number, (values, merged_cols) in data_rows():
        # A row with only values propagated from a merged range has no data of its own.
        if not _own_values(values, merged_cols):
            pending_empty += 1
            continue
        result.skipped_empty_rows += pending_empty
        pending_empty = 0
        if max_rows is not None and len(result.rows) >= max_rows:
            raise IngestionFailed(
                f"El archivo supera el límite de {max_rows} filas de datos por lote; "
                "divídalo en varios archivos.",
                code="row_limit_exceeded",
                details={"max_rows": max_rows},
            )
        raw: dict[str, Any] = {}
        merged_names: list[str] = []
        for index, value in enumerate(values, start=1):
            if index > len(headers):
                if value is None:
                    continue
                while len(headers) < index:  # data beyond the last header gets a generated name
                    headers.append(_unique(_generated_name(len(headers) + 1), seen))
            raw[headers[index - 1]] = value
            if index in merged_cols:
                merged_names.append(headers[index - 1])
        for name in headers:
            raw.setdefault(name, None)
        if merged_names:
            raw[MERGED_KEY] = merged_names
        result.rows.append(IngestedRow(row_number=row_number, raw_data=raw))
    return result
