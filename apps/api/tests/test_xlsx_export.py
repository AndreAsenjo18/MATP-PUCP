"""Tests for streaming Excel generator with bounded memory (Task 2.1, RF-036)."""

import datetime
import tracemalloc
from pathlib import Path

import openpyxl

from app.core.xlsx import write_table
from app.modules.search.filters import ActiveFilter


def test_write_table_creates_expected_sheets_and_content(tmp_path: Path) -> None:
    """Workbook has 'Resultados' with tabular data and 'Filtros' with audit metadata."""
    out_file = tmp_path / "test_export.xlsx"
    columns = ["ID", "Denominacin", "Coleccin", "Ao"]
    rows = [
        ["uuid-1", "Retablo Ayacuchano", "Coleccin RA", 1975],
        ["uuid-2", "Mate Burilado", "Coleccin Tradicional", 1980],
    ]
    filters = [
        ActiveFilter(key="collection_code", label="Coleccin", value_display="RA"),
        ActiveFilter(key="tenure_regime", label="Rgimen", value_display="Propiedad"),
    ]
    omitted = ["lender_name", "valuation"]

    write_table(
        out_file,
        rows,
        columns,
        sheet_title="Resultados",
        active_filters=filters,
        generated_at=datetime.datetime(2026, 10, 22, 10, 0, 0, tzinfo=datetime.UTC),
        user_label="catalogador@museo.gob.pe",
        omitted_columns=omitted,
    )

    wb = openpyxl.load_workbook(out_file)
    assert wb.sheetnames == ["Resultados", "Filtros"]

    # Check Resultados sheet
    ws_res = wb["Resultados"]
    res_rows = list(ws_res.iter_rows(values_only=True))
    assert res_rows[0] == ("ID", "Denominacin", "Coleccin", "Ao")
    assert res_rows[1] == ("uuid-1", "Retablo Ayacuchano", "Coleccin RA", 1975)
    assert res_rows[2] == ("uuid-2", "Mate Burilado", "Coleccin Tradicional", 1980)

    # Check Filtros sheet
    ws_filt = wb["Filtros"]
    filt_rows = list(ws_filt.iter_rows(values_only=True))
    assert ("Metadato", "Valor") in filt_rows
    assert any(
        r[0] == "Usuario solicitante" and r[1] == "catalogador@museo.gob.pe" for r in filt_rows
    )
    assert any(
        "lender_name, valuation" in str(r[1]) for r in filt_rows if r[0] and "omitidas" in r[0]
    )
    assert any(r[0] == "Coleccin" and r[1] == "RA" for r in filt_rows)
    assert any(r[0] == "Rgimen" and r[1] == "Propiedad" for r in filt_rows)


def test_write_table_bounded_memory_50k_rows(tmp_path: Path) -> None:
    """Streaming 50 000 rows with write_only=True uses bounded constant memory (< 35 MB peak)."""
    out_file = tmp_path / "large_50k.xlsx"
    columns = ["Cdigo", "Denominacin", "Autor", "Procedencia", "Ao"]

    def generate_50k():
        for i in range(50_000):
            yield [
                f"I-{i:05d}",
                f"Pieza de coleccin nmero {i}",
                f"Autor de prueba {i % 100}",
                f"Regin {i % 25}",
                1900 + (i % 120),
            ]

    tracemalloc.start()
    tracemalloc.reset_peak()

    write_table(
        out_file,
        generate_50k(),
        columns,
        user_label="admin@museo.gob.pe",
    )

    _current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    peak_mb = peak / (1024 * 1024)
    print(f"50,000 rows Excel write peak memory: {peak_mb:.2f} MB")

    # Peak memory must remain strictly bounded (< 35 MB)
    assert peak_mb < 35.0, f"Memory peaked at {peak_mb:.2f} MB, exceeding bounded threshold"
    assert out_file.exists()
    assert out_file.stat().st_size > 1_000_000  # Generated file is > 1 MB
